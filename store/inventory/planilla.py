# -*- coding: utf-8 -*-
"""Genera la planilla de emergencia (XLSX) para vender sin sistema."""
from io import BytesIO
from django.http import HttpResponse
from django.conf import settings
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from .models import Products

NAVY = "1A3C5E"
CUR = '"$" #,##0.00'
LINES = 12
VENTAS = 10


def _build_bytes():
    prods = list(Products.objects.filter(status=1).order_by('name'))
    rows = []
    for p in prods:
        rows.append(((p.codigo_barras or ''), p.name,
                     float(p.precio_minorista or 0), p.tipo_venta,
                     (int(p.plu) if p.plu else None)))
    rows.sort(key=lambda r: r[1].upper())

    navy = PatternFill("solid", fgColor=NAVY)
    inp = PatternFill("solid", fgColor="FBFCEE")
    totfill = PatternFill("solid", fgColor="FFE9A8")
    thin = Side(style='thin', color='B8C2CC')
    bd = Border(left=thin, right=thin, top=thin, bottom=thin)
    wbold = Font(name='Arial', bold=True, color='FFFFFF', size=10)
    ar = Font(name='Arial', size=10)
    arb = Font(name='Arial', size=10, bold=True)

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "VENTA"
    ws.sheet_view.showGridLines = False
    ws.merge_cells('A1:F1')
    ws['A1'] = "VENTAS DE EMERGENCIA - Lo de Angelito"
    ws['A1'].font = Font(name='Arial', size=15, bold=True, color=NAVY)
    ws.merge_cells('A2:F5')
    ws['A2'] = ("COMO USAR (sin internet ni sistema):\n"
                "- Con codigo de barras: escanealo en CODIGO. Descripcion y precio se completan solos. Pone la CANTIDAD.\n"
                "- Fraccionable (fiambre/queso): pesa en la balanza, escribi el PLU en la columna PLU (aparece la descripcion) "
                "y escribi en PRECIO el IMPORTE que muestra la balanza (deja CANT vacio).\n"
                "- Cada VENTA es una tabla: llenala, cobra el TOTAL, y la proxima venta va en la de abajo. "
                "Cuando vuelva luz/internet, se cargan en el sistema.")
    ws['A2'].font = Font(name='Arial', size=9)
    ws['A2'].alignment = Alignment(wrap_text=True, vertical='top')

    headers = ["CODIGO (escanea)", "PLU", "DESCRIPCION", "PRECIO", "CANT", "SUBTOTAL"]

    def block(start, num):
        ws.merge_cells(start_row=start, start_column=1, end_row=start, end_column=3)
        tc = ws.cell(row=start, column=1, value="VENTA #%d" % num)
        tc.font = Font(name='Arial', size=12, bold=True, color='FFFFFF'); tc.fill = navy
        tc.alignment = Alignment(horizontal='left', vertical='center', indent=1)
        ws.merge_cells(start_row=start, start_column=4, end_row=start, end_column=6)
        fc = ws.cell(row=start, column=4, value="Fecha/Hora: ________   Pago: ________")
        fc.font = Font(name='Arial', size=9, color='FFFFFF'); fc.fill = navy
        fc.alignment = Alignment(horizontal='right', vertical='center', indent=1)
        hr = start + 1
        for i, h in enumerate(headers, start=1):
            c = ws.cell(row=hr, column=i, value=h)
            c.font = wbold; c.fill = PatternFill("solid", fgColor="3B5872"); c.border = bd
            c.alignment = Alignment(horizontal='center', vertical='center')
        first = hr + 1
        last = first + LINES - 1
        for r in range(first, last + 1):
            ws.cell(row=r, column=1).number_format = '@'
            ws.cell(row=r, column=3, value=(
                '=IF($A%d<>"",IFERROR(VLOOKUP($A%d,PRECIOS!$A:$C,2,FALSE),"?"),'
                'IF($B%d<>"",IFERROR(INDEX(PRECIOS!$B:$B,MATCH($B%d,PRECIOS!$E:$E,0)),"?"),""))'
                % (r, r, r, r)))
            ws.cell(row=r, column=4, value='=IF($A%d<>"",IFERROR(VLOOKUP($A%d,PRECIOS!$A:$C,3,FALSE),0),"")' % (r, r))
            ws.cell(row=r, column=6, value='=IF($C%d="","",IF($D%d="",0,$D%d)*IF($E%d="",1,$E%d))' % (r, r, r, r, r))
            for col in range(1, 7):
                cc = ws.cell(row=r, column=col); cc.font = ar; cc.border = bd
                if col in (1, 2, 4, 5):
                    cc.fill = inp
            ws.cell(row=r, column=4).number_format = CUR
            ws.cell(row=r, column=6).number_format = CUR
        tr = last + 1
        ws.merge_cells(start_row=tr, start_column=1, end_row=tr, end_column=5)
        lc = ws.cell(row=tr, column=1, value="TOTAL VENTA #%d" % num)
        lc.font = arb; lc.alignment = Alignment(horizontal='right', indent=1)
        tot = ws.cell(row=tr, column=6, value='=SUM(F%d:F%d)' % (first, last))
        tot.font = Font(name='Arial', size=12, bold=True, color=NAVY); tot.number_format = CUR
        tot.fill = totfill; tot.border = bd
        return tr + 2

    r = 7
    for n in range(1, VENTAS + 1):
        r = block(r, n)
    for col, w in {'A': 17, 'B': 7, 'C': 40, 'D': 13, 'E': 8, 'F': 14}.items():
        ws.column_dimensions[col].width = w
    ws.freeze_panes = 'A7'

    wp = wb.create_sheet("PRECIOS")
    wp.sheet_view.showGridLines = False
    wp.merge_cells('A1:E1')
    wp['A1'] = "LISTA DE PRECIOS (referencia) - se actualiza sola"
    wp['A1'].font = Font(name='Arial', size=14, bold=True, color=NAVY)
    for i, h in enumerate(["CODIGO DE BARRAS", "DESCRIPCION", "PRECIO", "TIPO", "PLU"], start=1):
        c = wp.cell(row=2, column=i, value=h); c.font = wbold; c.fill = navy; c.border = bd
        c.alignment = Alignment(horizontal='center', vertical='center')
    rr = 3
    for barcode, name, precio, tipo, plu in rows:
        a = wp.cell(row=rr, column=1, value=barcode); a.number_format = '@'
        wp.cell(row=rr, column=2, value=name)
        pc = wp.cell(row=rr, column=3, value=precio); pc.number_format = CUR
        wp.cell(row=rr, column=4, value=("Fracc. (x kilo)" if tipo == 'fraccionable' else "Unidad"))
        wp.cell(row=rr, column=5, value=plu)
        for col in range(1, 6):
            cc = wp.cell(row=rr, column=col); cc.font = ar; cc.border = bd
            if rr % 2 == 0:
                cc.fill = PatternFill("solid", fgColor="EEF2F6")
        rr += 1
    for col, w in {'A': 20, 'B': 44, 'C': 14, 'D': 15, 'E': 8}.items():
        wp.column_dimensions[col].width = w
    wp.freeze_panes = 'A3'

    bio = BytesIO()
    wb.save(bio)
    bio.seek(0)
    return bio.read()


def planilla_emergencia(request):
    token = request.GET.get('token', '')
    valid = getattr(settings, 'PLANILLA_EMERGENCIA_TOKEN', None)
    if not (request.user.is_authenticated or (valid and token == valid)):
        return HttpResponse('No autorizado', status=403)
    data = _build_bytes()
    resp = HttpResponse(
        data,
        content_type='application/vnd.openxmlformats-officedocument.spreadsheetml.sheet')
    resp['Content-Disposition'] = 'attachment; filename="Ventas_emergencia.xlsx"'
    return resp
